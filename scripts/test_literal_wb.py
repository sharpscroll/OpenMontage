import openpyxl

path = r"\\?\E:\Data\YT\Video Data\CAUSE & MOTION - Simple Production Hub.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
print("Loaded:", wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value)
            if "Q61" in val or "q61" in val:
                print(f"Sheet '{name}' R{r}C{c}: {val[:120]}")
