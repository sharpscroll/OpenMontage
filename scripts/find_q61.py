import openpyxl

path = r"C:\Users\Hp\.gemini\antigravity\brain\3b239e0d-cf08-4230-b1b3-e372d8f5148c\.user_uploaded\media_1793086380082_3bd57b4c.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb['06 SHORTS BANK']

for r_idx in range(1, ws.max_row + 1):
    vals = [ws.cell(r_idx, c_idx).value for c_idx in range(1, 15)]
    r_str = " ".join([str(v) for v in vals if v is not None])
    if "61" in r_str or "Earth" in r_str or "Washer" in r_str:
        print(f"Row {r_idx}: {vals[:6]}")
