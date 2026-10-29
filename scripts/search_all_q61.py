import openpyxl

path = r"C:\Users\Hp\.gemini\antigravity\brain\3b239e0d-cf08-4230-b1b3-e372d8f5148c\.user_uploaded\media_1793086380082_3bd57b4c.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

for name in wb.sheetnames:
    ws = wb[name]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(r, c).value)
            if "Q61" in val or "q61" in val:
                print(f"Sheet '{name}' R{r}C{c}: {val[:120]}")
