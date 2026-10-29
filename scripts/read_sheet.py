import openpyxl

path = r"C:\Users\Hp\.gemini\antigravity\brain\3b239e0d-cf08-4230-b1b3-e372d8f5148c\.user_uploaded\media_1793086380082_3bd57b4c.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb['06 SHORTS BANK']

print(f"Max rows: {ws.max_row}, Max cols: {ws.max_column}")
for r_idx in range(1, min(ws.max_row + 1, 100)):
    vals = [ws.cell(r_idx, c_idx).value for c_idx in range(1, 12)]
    if any(vals):
        print(f"Row {r_idx:2d}: {vals}")
